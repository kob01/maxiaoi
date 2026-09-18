"""Multi-modal document parsers producing section-level blocks.

Each parser returns a list of ``ParsedBlock`` (section text + page number +
section heading), which the ingestion pipeline turns into parent/child
chunks. Supported types:

- text:  txt / md (md split by headings) / pdf (per real page) / docx
         (split by Heading styles) / pptx (per slide) / xlsx (per sheet)
- video transcripts: srt / vtt / *.transcript.txt (timeline merged)
- images: jpg / jpeg / png / webp / bmp -> captioned by the local vision
  model (qwen3-vl via Ollama) so image content becomes searchable text.
"""

from __future__ import annotations

import base64
import re
from dataclasses import dataclass
from pathlib import Path

import httpx

from app.config import get_settings

SUPPORTED_TEXT_EXT = {".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx"}
SUPPORTED_VIDEO_EXT = (".srt", ".vtt", ".transcript.txt")
SUPPORTED_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}

IMAGE_CAPTION_PROMPT = (
    "请详细描述这张图片中的所有文字内容、图表结构与关键业务信息，"
    "输出连贯的中文纯文本，用于企业知识库检索。"
)


@dataclass
class ParsedBlock:
    """A section-level block of parsed document text."""

    section: str       # heading path / slide title / sheet name; "" if none
    page_no: int       # pdf page / pptx slide / xlsx sheet number; -1 unknown
    text: str


def supported_extensions() -> set[str]:
    """All uploadable extensions."""
    return SUPPORTED_TEXT_EXT | set(SUPPORTED_VIDEO_EXT) | SUPPORTED_IMAGE_EXT


def modality_of(path: Path) -> str:
    """Map a file to its modality: text / video_transcript / image."""
    name = path.name.lower()
    if any(name.endswith(ext) for ext in SUPPORTED_VIDEO_EXT):
        return "video_transcript"
    if path.suffix.lower() in SUPPORTED_IMAGE_EXT:
        return "image"
    return "text"


# ---------------- text-family parsers ----------------


def _parse_txt(path: Path) -> list[ParsedBlock]:
    text = path.read_text(encoding="utf-8")
    return [ParsedBlock(section="", page_no=-1, text=text)] if text.strip() else []


def _parse_md(path: Path) -> list[ParsedBlock]:
    """Split markdown by headings; section = current heading path."""
    text = path.read_text(encoding="utf-8")
    blocks: list[ParsedBlock] = []
    heading_stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            section = "/".join(h for _, h in heading_stack)
            blocks.append(ParsedBlock(section=section, page_no=-1, text=body))
        buf.clear()

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            level = len(m.group(1))
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            heading_stack.append((level, m.group(2).strip()))
            buf.append(line)
        else:
            buf.append(line)
    flush()
    return blocks


def _parse_pdf(path: Path) -> list[ParsedBlock]:
    """One block per page so chunks keep the real page number."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[ParsedBlock] = []
    for idx, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(ParsedBlock(section="", page_no=idx, text=text))
    return blocks


def _parse_docx(path: Path) -> list[ParsedBlock]:
    """Split docx by Heading styles (supports EN/CN style names)."""
    import docx

    document = docx.Document(str(path))
    blocks: list[ParsedBlock] = []
    heading_stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            section = "/".join(h for _, h in heading_stack)
            blocks.append(ParsedBlock(section=section, page_no=-1, text=body))
        buf.clear()

    for para in document.paragraphs:
        style = para.style.name if para.style else ""
        m = re.match(r"^(?:Heading|标题)\s*(\d+)$", style)
        if m and para.text.strip():
            flush()
            level = int(m.group(1))
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            heading_stack.append((level, para.text.strip()))
            buf.append(para.text.strip())
        elif para.text.strip():
            buf.append(para.text.strip())
    flush()
    return blocks


def _parse_pptx(path: Path) -> list[ParsedBlock]:
    """One block per slide; section = slide title, page_no = slide number."""
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[ParsedBlock] = []
    for idx, slide in enumerate(prs.slides, 1):
        title = ""
        texts: list[str] = []
        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip()
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in p.runs).strip()
                    if line:
                        texts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells)
                    if line.strip(" |"):
                        texts.append(line)
        body = "\n".join(texts).strip()
        if body:
            blocks.append(ParsedBlock(section=title or f"第{idx}页", page_no=idx, text=body))
    return blocks


def _parse_xlsx(path: Path) -> list[ParsedBlock]:
    """One block per sheet; rows rendered as `cell | cell` lines."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[ParsedBlock] = []
    for idx, sheet in enumerate(wb.worksheets, 1):
        lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        body = "\n".join(lines).strip()
        if body:
            blocks.append(ParsedBlock(section=sheet.title, page_no=idx, text=body))
    wb.close()
    return blocks


def _parse_subtitle(path: Path) -> list[ParsedBlock]:
    """Parse srt/vtt (or plain transcript) into one timeline-annotated block."""
    raw = path.read_text(encoding="utf-8").replace("WEBVTT", "")
    block_re = re.compile(
        r"(?P<start>\d{2}:\d{2}:\d{2}[,.]\d{3})\s*-->\s*"
        r"(?P<end>\d{2}:\d{2}:\d{2}[,.]\d{3})\s*\n(?P<text>.*?)(?=\n\s*\n|\Z)",
        re.DOTALL,
    )
    lines: list[str] = []
    for m in block_re.finditer(raw):
        text = re.sub(r"<[^>]+>", "", m.group("text")).strip()
        if text:
            lines.append(f"[{m.group('start')} -> {m.group('end')}] {text}")
    if not lines and raw.strip():  # plain transcript fallback
        lines.append(raw.strip())
    return [ParsedBlock(section="", page_no=-1, text="\n".join(lines))] if lines else []


# ---------------- image parser (vision model) ----------------

# 1568px preserves OCR readability while keeping the request fast.
_IMAGE_MAX_EDGE = 1568


def _compress_image(path: Path) -> bytes:
    """Load, flatten to RGB, downscale to _IMAGE_MAX_EDGE and re-encode JPEG."""
    from io import BytesIO

    from PIL import Image

    with Image.open(path) as img:
        img = img.convert("RGB")
        w, h = img.size
        longest = max(w, h)
        if longest > _IMAGE_MAX_EDGE:
            scale = _IMAGE_MAX_EDGE / longest
            img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))))
        buf = BytesIO()
        img.save(buf, "JPEG", quality=85)
        return buf.getvalue()


async def _parse_image(path: Path) -> list[ParsedBlock]:
    """Caption an image with the local vision model (qwen3-vl)."""
    settings = get_settings()
    try:
        data = _compress_image(path)
    except Exception as exc:
        raise RuntimeError(f"图片文件无法读取或已损坏: {path.name} ({exc})") from exc
    b64 = base64.b64encode(data).decode()
    payload = {
        "model": settings.vision_model,
        "prompt": IMAGE_CAPTION_PROMPT,
        "images": [b64],
        "stream": False,
    }
    async with httpx.AsyncClient(timeout=float(settings.vision_timeout)) as client:
        try:
            resp = await client.post(
                f"{settings.ollama_base_url.rstrip('/')}/api/generate", json=payload
            )
            resp.raise_for_status()
        except httpx.TimeoutException as exc:
            raise RuntimeError(
                f"图片解析超时(视觉模型 {settings.vision_model}, 阈值 "
                f"{settings.vision_timeout}s): 冷启动较慢或图片过大, 可在 .env 调大 VISION_TIMEOUT 后重试"
            ) from exc
        except Exception as exc:
            raise RuntimeError(
                f"图片解析失败(视觉模型 {settings.vision_model} 不可用, "
                f"请确认已执行 ollama pull {settings.vision_model}): {exc}"
            ) from exc
    text = (resp.json().get("response") or "").strip()
    return [ParsedBlock(section="图片内容", page_no=-1, text=text)] if text else []


# ---------------- unified entry ----------------


async def parse_blocks(path: Path) -> tuple[str, list[ParsedBlock]]:
    """Parse any supported file into (modality, section blocks)."""
    modality = modality_of(path)
    name = path.name.lower()
    if modality == "video_transcript":
        return modality, _parse_subtitle(path)
    if modality == "image":
        return modality, await _parse_image(path)

    ext = path.suffix.lower()
    if name.endswith(".transcript.txt"):
        return "video_transcript", _parse_subtitle(path)
    if ext == ".txt":
        return modality, _parse_txt(path)
    if ext == ".md":
        return modality, _parse_md(path)
    if ext == ".pdf":
        return modality, _parse_pdf(path)
    if ext == ".docx":
        return modality, _parse_docx(path)
    if ext == ".pptx":
        return modality, _parse_pptx(path)
    if ext == ".xlsx":
        return modality, _parse_xlsx(path)
    raise ValueError(f"Unsupported file type: {path.name}")
